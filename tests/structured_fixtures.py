"""Unmistakably fictional documents generated in memory for MVP-04 tests."""

from __future__ import annotations

import io


def figure_png() -> bytes:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (160, 70), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((10, 10, 150, 60), outline="navy", width=3)
    draw.line((25, 45, 70, 22, 135, 30), fill="navy", width=3)
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def pdf_bytes(*, scanned: bool = False) -> bytes:
    import pymupdf

    document = pymupdf.open()
    page = document.new_page(width=600, height=790)
    if scanned:
        page.insert_image(pymupdf.Rect(40, 80, 500, 280), stream=figure_png())
    else:
        page.insert_text((50, 70), "Synthetic Battery Results", fontsize=17)
        page.insert_text((50, 110), "Capacity changed by -3.2 mAh g-1 at 25 C.", fontsize=11)
        page.insert_text((50, 145), "Figure 1. Fictional cycling curve.", fontsize=11)
        page.insert_image(pymupdf.Rect(50, 160, 210, 230), stream=figure_png())
        left, top, width, height = 50, 280, 420, 35
        for x in (left, left + 170, left + 300, left + width):
            page.draw_line((x, top), (x, top + 2 * height), color=(0, 0, 0))
        for y in (top, top + height, top + 2 * height):
            page.draw_line((left, y), (left + width, y), color=(0, 0, 0))
        for y, values in (
            (top + 22, ("Metric", "Value", "Unit")),
            (top + height + 22, ("Capacity", "-3.2", "mAh g-1")),
        ):
            for x, value in zip((left + 8, left + 178, left + 308), values):
                page.insert_text((x, y), value, fontsize=11)
    content = document.tobytes()
    document.close()
    return content


def docx_bytes() -> bytes:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_heading("Synthetic Electrochemistry", level=1)
    document.add_paragraph("The fictional voltage was -0.25 V vs reference; x² is symbolic.")
    table = document.add_table(rows=2, cols=2)
    for cell, text in zip(table.rows[0].cells, ("Condition", "Value (mAh g-1)")):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ("Cycle 1", "-3.2")):
        cell.text = text
    document.add_picture(io.BytesIO(figure_png()), width=Inches(1.5))
    document.add_paragraph("Figure 1. Fictional electrode image.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
    box.text = "Synthetic slide: -3.2 mAh g-1"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Unit"
    table.cell(1, 0).text = "Capacity"
    table.cell(1, 1).text = "mAh g-1"
    slide.shapes.add_picture(io.BytesIO(figure_png()), Inches(1), Inches(4))
    slide.notes_slide.notes_text_frame.text = "Speaker note: fictional measurement only."
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Measurements"
    sheet["A1"] = "Capacity (mAh g-1)"
    sheet["B1"] = "Adjusted"
    sheet["A2"] = -3.2
    sheet["B2"] = "=A2*2"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
