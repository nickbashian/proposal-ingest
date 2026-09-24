"""MVP-04 structured extraction, provenance, failure, and inspector contracts."""

from __future__ import annotations

from pathlib import Path
import os
import hashlib
import subprocess
import io
import re
import zipfile
import json

import pytest
from django.contrib.auth import get_user_model
from django.test import Client
from django.urls import reverse
from django.core.management import call_command
from django.db import connection
from django.test.utils import CaptureQueriesContext
from playwright.sync_api import sync_playwright

from scripts import dev

from proposal_app import extraction_audit, extraction_service, models as m, services, workflow
from proposal_app.structured_extraction import extract_snapshot
from tests.structured_fixtures import docx_bytes, pdf_bytes, pptx_bytes, xlsx_bytes

pytestmark = pytest.mark.django_db


@pytest.fixture
def owner(settings, tmp_path):
    settings.LOCAL_STORAGE_ROOT = tmp_path / "objects"
    user = get_user_model().objects.create_user(username="mvp04-owner")
    m.Identity.objects.create(user=user, issuer="local", subject="mvp04-owner", allowed=True)
    collection = m.Collection.objects.create(name="Fictional extraction collection")
    m.CollectionAccess.objects.create(user=user, collection=collection)
    return user, collection


def capture(owner, name: str, content: bytes) -> m.SourceVersion:
    user, collection = owner
    return services.observe_source(
        user,
        collection.id,
        identity={
            "connector": "local",
            "tenant": "local",
            "site": "local",
            "drive": "fixture",
            "item": name,
        },
        path=f"2025/Fictional Proposal/{name}",
        observation_key="initial",
        content=content,
        proposal="Fictional Proposal",
        family="source",
    )


def test_pdf_structure_tables_figures_and_exact_characters(settings):
    result = extract_snapshot("fictional.pdf", pdf_bytes(), settings.APP)
    assert result.state == "succeeded"
    assert any(unit.kind == "heading" and unit.locator["page"] == 1 for unit in result.units)
    assert any("-3.2 mAh g-1" in unit.text for unit in result.units)
    rows = [unit for unit in result.units if unit.kind == "table_row"]
    assert rows[0].context["headers"] == ["Metric", "Value", "Unit"]
    assert "Capacity | -3.2 | mAh g-1" in rows[1].text
    assert any(unit.kind == "caption" and "Figure 1" in unit.text for unit in result.units)
    assert len(result.figures) == 1
    assert result.figures[0].locator["page"] == 1
    assert result.figures[0].caption == "Figure 1. Fictional cycling curve."
    assert [unit.kind for unit in result.units[:3]] == ["heading", "paragraph", "caption"]


def test_docx_slide_and_spreadsheet_locators(settings):
    document = extract_snapshot("fictional.docx", docx_bytes(), settings.APP)
    assert document.state == "succeeded"
    assert any(
        unit.kind == "heading" and unit.locator["section"] == "Synthetic Electrochemistry"
        for unit in document.units
    )
    assert any("-0.25 V" in unit.text and "x²" in unit.text for unit in document.units)
    assert any(
        unit.kind == "table_cell" and unit.locator["cell"] == 2 and "mAh g-1" in unit.text
        for unit in document.units
    )
    assert all("page" not in unit.locator for unit in document.units)
    assert len(document.figures) == 1
    assert document.figures[0].caption == "Figure 1. Fictional electrode image."

    slides = extract_snapshot("fictional.pptx", pptx_bytes(), settings.APP)
    assert slides.state == "succeeded"
    assert any(unit.locator.get("slide") == 1 and "-3.2" in unit.text for unit in slides.units)
    assert any(unit.kind == "notes" and unit.locator["notes"] for unit in slides.units)
    assert any(unit.kind == "table_cell" for unit in slides.units)
    assert len(slides.figures) == 1

    spreadsheet = extract_snapshot("fictional.xlsx", xlsx_bytes(), settings.APP)
    assert spreadsheet.state == "succeeded"
    assert any(
        unit.locator["sheet"] == "Measurements" and "A2" in unit.locator["cells"]
        for unit in spreadsheet.units
    )
    formula = next(unit for unit in spreadsheet.units if "=A2*2" in unit.text)
    assert formula.context["formula_cells"][0]["cached_value"] is None
    assert "formula_cached_value_not_recalculated" in formula.warnings


@pytest.mark.parametrize(
    ("name", "content_kind", "reason"),
    [
        ("scan.pdf", "scan", "no_selectable_text"),
        ("empty.txt", "empty", "empty_file"),
        ("broken.pdf", "broken", "malformed_or_unreadable"),
        ("old.doc", "legacy", "legacy_format"),
        ("archive.zip", "archive", "unsupported_format"),
    ],
)
def test_unreadable_inputs_have_recovery(settings, name, content_kind, reason):
    content = {
        "scan": lambda: pdf_bytes(scanned=True),
        "empty": lambda: b"",
        "broken": lambda: b"not a PDF",
        "legacy": lambda: b"legacy",
        "archive": lambda: b"archive",
    }[content_kind]()
    result = extract_snapshot(name, content, settings.APP)
    assert result.state != "succeeded"
    assert result.reason == reason
    assert result.recovery_action
    assert not result.units


def test_encrypted_pdf_and_resource_limits(settings):
    import pymupdf

    document = pymupdf.open(stream=pdf_bytes(), filetype="pdf")
    encrypted = document.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="fictional-owner", user_pw="fictional-user"
    )
    document.close()
    assert extract_snapshot("locked.pdf", encrypted, settings.APP).reason == "encrypted_pdf"
    limits = {**settings.APP, "extraction_max_source_bytes": 10}
    assert extract_snapshot("too-big.pdf", pdf_bytes(), limits).reason == "source_size_limit"
    limits = {**settings.APP, "extraction_max_pages": 0}
    assert extract_snapshot("page.pdf", pdf_bytes(), limits).reason == "page_limit"
    limits = {**settings.APP, "extraction_max_spreadsheet_cells": 1}
    assert extract_snapshot("sheet.xlsx", xlsx_bytes(), limits).reason == "spreadsheet_cell_limit"


def test_cache_reextraction_read_only_and_historical_citation(owner, settings, tmp_path):
    user, collection = owner
    source = tmp_path / "fictional.pdf"
    source.write_bytes(pdf_bytes())
    original = source.read_bytes()
    original_mtime = source.stat().st_mtime_ns
    version = capture(owner, source.name, original)
    first = extraction_service.extract_version(user, version.id)
    assert first.state == "succeeded" and first.active
    assert extraction_service.extract_version(user, version.id).id == first.id
    first_unit = m.ExtractedUnit.objects.filter(extraction_run=first).first()
    assert first_unit
    family = m.VersionFamily.objects.get(proposal__collection=collection)
    decision = m.Decision.objects.create(
        family=family,
        scope=f"version:{version.id}",
        field="publication",
        kind="inclusion",
        revision=1,
    )
    event = m.DecisionEvent.objects.create(
        decision=decision,
        revision=1,
        actor=user,
        value={"treatment": "include"},
        rationale="Fictional test approval",
        evidence=[str(first_unit.id)],
        resolver_revision="test",
    )
    generation = m.PublicationGeneration.objects.create(
        proposal=family.proposal, revision=1, state="active"
    )
    artifact = m.PublicationArtifact.objects.create(
        generation=generation,
        unit=first_unit,
        decision_event=event,
        blob=version.blob,
        eligible=True,
    )
    client = Client()
    client.force_login(user)
    before = client.get(reverse("artifact", args=[artifact.id]))
    assert first_unit.text in before.content.decode()
    second = extraction_service.extract_version(user, version.id, force=True)
    assert second.id != first.id and second.active
    first.refresh_from_db()
    assert not first.active
    assert m.ExtractedUnit.objects.get(pk=first_unit.id).text == first_unit.text
    after = client.get(reverse("artifact", args=[artifact.id]))
    assert first_unit.text in after.content.decode()
    assert str(first_unit.id) == str(artifact.unit_id)
    assert source.read_bytes() == original and source.stat().st_mtime_ns == original_mtime
    assert m.ExtractionRun.objects.filter(version=version).count() == 2
    assert extraction_service.extraction_fingerprint(version) == first.fingerprint
    settings.APP = {**settings.APP, "extraction_revision": "structured-v2"}
    third = extraction_service.extract_version(user, version.id)
    assert third.id != second.id
    assert third.fingerprint != first.fingerprint
    assert m.ExtractedUnit.objects.get(pk=first_unit.id).text == first_unit.text


def test_inspector_auth_figures_selective_interpretation_and_failure(owner):
    user, collection = owner
    version = capture(owner, "figure.pdf", pdf_bytes())
    scanned = capture(owner, "scanned.pdf", pdf_bytes(scanned=True))
    client = Client()
    assert client.get(reverse("source-inspector", args=[version.id])).status_code == 401
    client.force_login(user)
    response = client.post(reverse("source-inspector", args=[version.id]), {"action": "extract"})
    assert response.status_code == 302
    run = m.ExtractionRun.objects.get(version=version)
    page = client.get(reverse("source-inspector", args=[version.id]))
    assert b"Original extracted units" in page.content
    assert b"pymupdf" in page.content
    assert b"layout_order_inferred" in page.content
    figure = m.FigureAsset.objects.get(run=run)
    image = client.get(reverse("figure-image", args=[figure.id]))
    assert image.status_code == 200 and image["Content-Type"] == "image/png"
    original = client.get(reverse("source-original", args=[version.id]))
    assert original.status_code == 200
    assert hashlib.sha256(original.content).hexdigest() == version.blob.sha256
    mismatched = client.post(
        reverse("source-inspector", args=[scanned.id]),
        {
            "action": "request-visual",
            "run_id": str(run.id),
            "kind": "figure",
            "figure_id": str(figure.id),
        },
    )
    assert mismatched.status_code == 404
    task = extraction_service.request_visual(
        user, run.id, kind="figure", locator=figure.locator, figure_id=figure.id
    )
    assert task.state == "requested"
    extraction_service.complete_visual(
        user,
        task.id,
        text="A fictional curve rises then falls.",
        source_check="Reviewed the preserved image on page 1; values were not digitized.",
    )
    assert b"Derived content" in client.get(reverse("source-inspector", args=[version.id])).content
    assert not m.ExtractedUnit.objects.filter(text__contains="rises then falls").exists()
    scanned_run = extraction_service.extract_version(user, scanned.id)
    assert scanned_run.state == "scanned"
    assert (
        b"no_selectable_text" in client.get(reverse("source-inspector", args=[scanned.id])).content
    )
    ocr_task = extraction_service.request_visual(
        user, scanned_run.id, kind="ocr", locator={"page": 1}
    )
    assert ocr_task.state == "requested"
    assert m.ExtractedUnit.objects.filter(version=scanned).count() == 0
    other = get_user_model().objects.create_user(username="mvp04-other")
    m.Identity.objects.create(user=other, issuer="local", subject="mvp04-other", allowed=True)
    client.force_login(other)
    assert client.get(reverse("source-inspector", args=[version.id])).status_code == 403
    assert client.get(reverse("figure-image", args=[figure.id])).status_code == 403
    assert client.get(reverse("source-original", args=[version.id])).status_code == 403


def test_figure_only_office_has_no_pdf_ocr_and_tiff_has_preview(owner, monkeypatch):
    from docx import Document
    from docx.shared import Inches
    from PIL import Image

    user, _ = owner
    picture = io.BytesIO()
    Image.new("RGB", (20, 10), "white").save(picture, format="TIFF")
    document = Document()
    document.add_picture(io.BytesIO(picture.getvalue()), width=Inches(1))
    saved = io.BytesIO()
    document.save(saved)
    version = capture(owner, "picture.docx", saved.getvalue())
    run = extraction_service.extract_version(user, version.id)
    assert run.state == "failed" and run.reason == "figure_only_non_pdf"
    figure = m.FigureAsset.objects.get(run=run)
    assert figure.mime_type == "image/tiff"
    with pytest.raises(ValueError, match="PDF"):
        extraction_service.request_visual(user, run.id, kind="ocr", locator={"page": 1})
    client = Client()
    client.force_login(user)
    original = client.get(reverse("figure-image", args=[figure.id]))
    preview = client.get(reverse("figure-image", args=[figure.id]) + "?preview=1")
    assert original["Content-Type"] == "image/tiff"
    assert original.content == picture.getvalue()
    assert preview["Content-Type"] == "image/png"
    assert preview.content.startswith(b"\x89PNG")

    def oversized_image(*args, **kwargs):
        raise Image.DecompressionBombError()

    monkeypatch.setattr(Image, "open", oversized_image)
    assert client.get(reverse("figure-image", args=[figure.id]) + "?preview=1").status_code == 404


def test_inspector_post_keeps_selected_historical_run(owner):
    user, _ = owner
    version = capture(owner, "history.pdf", pdf_bytes())
    first = extraction_service.extract_version(user, version.id)
    extraction_service.extract_version(user, version.id, force=True)
    figure = m.FigureAsset.objects.get(run=first)
    client = Client()
    client.force_login(user)
    response = client.post(
        reverse("source-inspector", args=[version.id]) + f"?run={first.id}",
        {
            "action": "request-visual",
            "run_id": str(first.id),
            "kind": "figure",
            "figure_id": str(figure.id),
        },
    )
    assert response.status_code == 302
    assert response["Location"].endswith(f"?run={first.id}")


def test_failed_extraction_is_visible_and_not_active(owner):
    user, _ = owner
    version = capture(owner, "malformed.pdf", b"not a PDF")
    run = extraction_service.extract_version(user, version.id)
    assert run.state == "failed" and not run.active
    assert run.reason == "malformed_or_unreadable"
    assert not m.ExtractedUnit.objects.filter(version=version).exists()
    assert extraction_service.extract_version(user, version.id).id == run.id


def test_source_path_at_capture_survives_rename(owner):
    user, _ = owner
    version = capture(owner, "original.pdf", pdf_bytes())
    version.source.display_path = "2025/Fictional Proposal/renamed.docx"
    version.source.save(update_fields=["display_path"])
    assert version.observed_path.endswith("original.pdf")
    run = extraction_service.extract_version(user, version.id)
    assert run.state == "succeeded"
    assert run.parser.startswith("pymupdf")


def test_raw_extracted_units_cannot_enter_factual_review(owner):
    user, collection = owner
    version = capture(owner, "review.pdf", pdf_bytes())
    first = extraction_service.extract_version(user, version.id)
    assert set(
        m.ExtractedUnit.objects.filter(extraction_run=first).values_list("support_kind", flat=True)
    ) == {"unclassified"}
    family = m.VersionFamily.objects.get(proposal__collection=collection)
    decision = m.Decision.objects.create(
        family=family, scope=f"version:{version.id}", field="publication", kind="inclusion"
    )
    with pytest.raises(ValueError, match="requires factual"):
        workflow.answer_inclusion(user, decision.id, 0, "include")


def test_active_run_controls_review_and_prior_approval_cannot_republish(owner):
    user, collection = owner
    version = capture(owner, "review.txt", b"Fictional evidence")
    family = m.VersionFamily.objects.get(proposal__collection=collection)
    decision = m.Decision.objects.create(
        family=family, scope=f"version:{version.id}", field="publication", kind="inclusion"
    )
    first = m.ExtractionRun.objects.create(
        version=version,
        number=1,
        fingerprint="1" * 64,
        extractor_revision="fictional",
        parser="fictional",
        state="succeeded",
        active=True,
    )
    old_unit = m.ExtractedUnit.objects.create(
        version=version,
        extraction_run=first,
        extractor_revision="fictional-1",
        key="fact",
        locator={"line": 1},
        text="Fictional old evidence",
        support_kind="factual",
    )
    old_event = workflow.answer_inclusion(user, decision.id, 0, "include")
    first.active = False
    first.save(update_fields=["active"])
    second = m.ExtractionRun.objects.create(
        version=version,
        number=2,
        fingerprint="2" * 64,
        extractor_revision="fictional",
        parser="fictional",
        state="succeeded",
        active=True,
    )
    new_unit = m.ExtractedUnit.objects.create(
        version=version,
        extraction_run=second,
        extractor_revision="fictional-2",
        key="fact",
        locator={"line": 1},
        text="Fictional updated evidence",
        support_kind="factual",
    )
    with pytest.raises(ValueError, match="active extraction"):
        workflow.publish(user, family.proposal_id)
    new_event = workflow.answer_inclusion(user, decision.id, 1, "include")
    assert new_event.evidence == [str(new_unit.id)]
    assert old_event.evidence == [str(old_unit.id)]


def test_transient_failure_retries_and_cached_success_reactivates(owner, settings, monkeypatch):
    user, _ = owner
    version = capture(owner, "retry.pdf", pdf_bytes())
    original = extraction_service._parse_bounded
    monkeypatch.setattr(
        extraction_service,
        "_parse_bounded",
        lambda *args: extraction_service.ExtractionResult(
            "failed", "parser_timeout", "Retry", "worker"
        ),
    )
    failed = extraction_service.extract_version(user, version.id)
    monkeypatch.setattr(extraction_service, "_parse_bounded", original)
    first = extraction_service.extract_version(user, version.id)
    assert first.id != failed.id and first.active
    settings.APP = {**settings.APP, "extraction_revision": "new-parser-revision"}
    second = extraction_service.extract_version(user, version.id)
    assert second.active
    settings.APP = {**settings.APP, "extraction_revision": "structured-v1"}
    reused = extraction_service.extract_version(user, version.id)
    assert reused.id == first.id and reused.active
    second.refresh_from_db()
    assert not second.active


def test_spreadsheet_blank_first_cell_and_missing_dimension(settings):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet["B2"] = "Fictional value"
    output = io.BytesIO()
    workbook.save(output)
    rebuilt = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(output.getvalue())) as original,
        zipfile.ZipFile(rebuilt, "w") as target,
    ):
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<dimension[^>]*/>", b"", data)
            target.writestr(item, data)
    result = extract_snapshot("blank.xlsx", rebuilt.getvalue(), settings.APP)
    assert result.state == "succeeded"
    assert result.units[0].locator["row"] == 2
    assert result.units[0].locator["cells"] == ["B2"]


def test_cache_fingerprint_ignores_visual_only_settings(owner, settings):
    version = capture(owner, "fingerprint.pdf", pdf_bytes())
    baseline = extraction_service.extraction_fingerprint(version)
    settings.APP = {
        **settings.APP,
        "extraction_ocr_executable": "fictional-ocr",
        "extraction_max_visual_tasks_per_version": 7,
    }
    assert extraction_service.extraction_fingerprint(version) == baseline
    settings.APP = {
        **settings.APP,
        "extraction_max_units": settings.APP["extraction_max_units"] + 1,
    }
    assert extraction_service.extraction_fingerprint(version) != baseline


def test_parser_package_revision_invalidates_extraction_cache(owner, monkeypatch):
    version = capture(owner, "library.pdf", pdf_bytes())
    baseline = extraction_service.extraction_fingerprint(version)
    monkeypatch.setattr(extraction_service, "parser_library_revision", lambda name: "PyMuPDF:new")
    assert extraction_service.extraction_fingerprint(version) != baseline


def test_audit_samples_failures_and_figures_without_source_text(owner):
    user, collection = owner
    readable = capture(owner, "readable.pdf", pdf_bytes())
    unreadable = capture(owner, "unreadable.pdf", b"malformed")
    extraction_service.extract_version(user, readable.id)
    extraction_service.extract_version(user, unreadable.id)
    report = extraction_audit.audit_collection(collection.id, per_family=2)
    family = report["families"][0]
    assert family["version_count"] == 2
    assert family["counts"]["failed"] == 1
    assert family["counts"]["figures"] == 1
    assert family["sample"][0]["source_version_id"] == str(unreadable.id)
    assert "Capacity" not in str(report)
    assert "Synthetic Battery Results" not in str(report)


def test_audit_counts_shared_versions_once_and_hides_extensionless_path(owner):
    user, collection = owner
    version = capture(owner, "private-synthetic-marker", b"fictional")
    extraction_service.extract_version(user, version.id)
    second = m.Proposal.objects.create(collection=collection, identifier="Second fictional family")
    family = m.VersionFamily.objects.create(proposal=second, key="source")
    m.ProposalMembership.objects.create(source=version.source, family=family)
    report = extraction_audit.audit_collection(collection.id)
    assert len(report["families"]) == 2
    assert report["totals"]["failed"] == 1
    assert all(row["format"] == "none" for item in report["families"] for row in item["sample"])
    output = io.StringIO()
    call_command("audit_extraction", collection=str(collection.id), stdout=output)
    assert "private-synthetic-marker" not in output.getvalue()


def test_audit_query_count_is_bounded_across_versions(owner):
    user, collection = owner
    for index in range(12):
        version = capture(owner, f"fictional-{index}.txt", b"Synthetic content")
        extraction_service.extract_version(user, version.id)
    with CaptureQueriesContext(connection) as queries:
        report = extraction_audit.audit_collection(collection.id)
    assert report["families"][0]["version_count"] == 12
    assert len(queries) < 12


def test_parser_worker_discards_output_and_receives_only_parser_settings(
    settings, tmp_path, monkeypatch
):
    settings.LOCAL_STORAGE_ROOT = tmp_path / "objects"

    def fake_worker(argv, **kwargs):
        assert kwargs["stdout"] == subprocess.DEVNULL
        assert kwargs["stderr"] == subprocess.DEVNULL
        config = json.loads(Path(argv[-1]).read_text(encoding="utf-8"))
        assert set(config) == set(extraction_service.PARSER_SETTING_KEYS)
        assert "development_database_url" not in config
        return subprocess.CompletedProcess(argv, 1)

    monkeypatch.setattr(extraction_service.subprocess, "run", fake_worker)
    result = extraction_service._parse_bounded("fictional.txt", b"Synthetic")
    assert result.reason == "parser_worker_error"


def test_selective_local_ocr_requires_source_check(owner, settings, monkeypatch):
    user, _ = owner
    version = capture(owner, "scanned-for-ocr.pdf", pdf_bytes(scanned=True))
    run = extraction_service.extract_version(user, version.id)
    task = extraction_service.request_visual(user, run.id, kind="ocr", locator={"page": 1})
    with pytest.raises(ValueError, match="unavailable"):
        extraction_service.run_selective_ocr(user, task.id)
    settings.APP = {**settings.APP, "extraction_ocr_executable": "fictional-ocr"}
    monkeypatch.setattr(extraction_service.shutil, "which", lambda name: name)

    class FakeOcr:
        returncode = 0

        def __init__(self, argv, **kwargs):
            self.argv = argv
            kwargs["stdout"].write(b"Possible -3.2 mAh g-1")

        def poll(self):
            return self.returncode

    def fake_ocr(argv, **kwargs):
        assert argv[0] == "fictional-ocr" and argv[2] == "stdout"
        assert Path(argv[1]).is_file()
        assert Path(argv[1]).read_bytes().startswith(b"\x89PNG")
        return FakeOcr(argv, **kwargs)

    monkeypatch.setattr(extraction_service.subprocess, "Popen", fake_ocr)
    candidate = extraction_service.run_selective_ocr(user, task.id)
    assert candidate.state == "needs_source_check"
    assert candidate.adapter_revision == "local-ocr-v1"
    assert not m.ExtractedUnit.objects.filter(version=version).exists()
    checked = extraction_service.complete_visual(
        user,
        task.id,
        text=candidate.interpretation,
        source_check="Compared visually with the selected page; minus sign uncertain.",
    )
    assert checked.state == "source_checked"
    assert not m.ExtractedUnit.objects.filter(version=version).exists()


def test_parser_timeout_is_recoverable(owner, monkeypatch):
    user, _ = owner
    version = capture(owner, "stalled.pdf", pdf_bytes())

    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired(args[0], 1)

    monkeypatch.setattr(extraction_service.subprocess, "run", timeout)
    run = extraction_service.extract_version(user, version.id)
    assert run.state == "failed" and run.reason == "parser_timeout"
    assert "alternate parser" in run.recovery_action
    assert not m.ExtractedUnit.objects.filter(version=version).exists()


@pytest.mark.django_db(transaction=True)
def test_browser_inspector_walkthrough(owner, settings, live_server, monkeypatch):
    _, collection = owner
    version = capture(owner, "fictional.pdf", pdf_bytes())
    document = capture(owner, "fictional.docx", docx_bytes())
    slides = capture(owner, "fictional.pptx", pptx_bytes())
    spreadsheet = capture(owner, "fictional.xlsx", xlsx_bytes())
    scanned = capture(owner, "scanned.pdf", pdf_bytes(scanned=True))
    malformed = capture(owner, "malformed.pdf", b"not a PDF")
    monkeypatch.setattr(settings, "LOCAL_AUTH", True)
    if not os.environ.get("PLAYWRIGHT_BROWSERS_PATH"):
        dev._configure_browser_cache()
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        try:
            page = browser.new_page()
            assert (
                page.goto(live_server.url + reverse("source-inspector", args=[version.id])).status
                == 401
            )
            page.goto(live_server.url + "/login/")
            page.get_by_label("Local identity").fill("mvp04-owner")
            page.get_by_role("button", name="Sign in locally").click()
            page.goto(live_server.url + reverse("collection", args=[collection.id]))
            page.get_by_role("row").filter(has_text="fictional.pdf").get_by_role(
                "link", name="Inspect source version 1"
            ).click()
            page.get_by_role("button", name="Extract or reuse cached result").click()
            body = page.locator("body").inner_text()
            assert "Capacity | -3.2 | mAh g-1" in body
            assert "Figure 1. Fictional cycling curve." in body
            assert "Preserved figures" in body
            page.get_by_role("button", name="Request selective figure interpretation").click()
            page.get_by_label("Derived interpretation or transcription").fill(
                "Fictional curve; no digitized values."
            )
            page.get_by_label("Source check and uncertainty").fill(
                "Compared against the preserved source image on page 1."
            )
            page.get_by_role("button", name="Save source-checked interpretation").click()
            assert (
                "Derived content, source checked by reviewer" in page.locator("body").inner_text()
            )
            if screenshot_path := os.environ.get("MVP04_SCREENSHOT_PATH"):
                destination = Path(screenshot_path)
                destination.parent.mkdir(parents=True, exist_ok=True)
                page.screenshot(path=destination, full_page=True)
            page.get_by_role("button", name="Re-extract as a new version").click()
            assert "Run #2: succeeded" in page.locator("body").inner_text()
            assert "#1 succeeded" in page.locator("body").inner_text()
            for target, expected in (
                (document, "x² is symbolic"),
                (slides, "Speaker note: fictional measurement only"),
                (spreadsheet, "cached: unavailable"),
            ):
                page.goto(live_server.url + reverse("source-inspector", args=[target.id]))
                page.get_by_role("button", name="Extract or reuse cached result").click()
                assert expected in page.locator("body").inner_text()
            page.goto(live_server.url + reverse("source-inspector", args=[scanned.id]))
            page.get_by_role("button", name="Extract or reuse cached result").click()
            assert "no_selectable_text" in page.locator("body").inner_text()
            assert "Request selective OCR" in page.locator("body").inner_text()
            page.goto(live_server.url + reverse("source-inspector", args=[malformed.id]))
            page.get_by_role("button", name="Extract or reuse cached result").click()
            assert "malformed_or_unreadable" in page.locator("body").inner_text()
        finally:
            browser.close()
    assert m.ExtractionRun.objects.count() == 7
